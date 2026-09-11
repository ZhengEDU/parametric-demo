"""
Shared helpers for building the three demo-authored templates (Temp/RH
Meter, Weathering Tester, Preventative Maintenance) from the same real
Parametric shell used for the Weight Set template (logo, header/footer,
fonts, legal boilerplate, signature block). See docs/ASSUMPTIONS.md —
only the Weight Set table structure is copied verbatim from an uploaded
real file; these three reuse the shell but the table/checklist bodies are
demo-authored from the written brief, not scanned.
"""
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "samples" / "word-documents" / "20260817.003.docx"


def load_shell():
    if not SRC.exists():
        raise SystemExit(f"Source sample not found: {SRC}")
    return docx.Document(str(SRC))


def move_after(anchor, tbl_or_para):
    """Detach tbl_or_para from wherever it currently is and place it
    immediately after `anchor` in document order. Returns the moved
    element so callers can chain further insertions after it."""
    el = tbl_or_para._tbl if hasattr(tbl_or_para, "_tbl") else tbl_or_para._p
    el.getparent().remove(el)
    anchor_el = anchor._tbl if hasattr(anchor, "_tbl") else anchor._p
    anchor_el.addnext(el)
    return tbl_or_para


def remove(el):
    xml_el = el._tbl if hasattr(el, "_tbl") else el._p
    xml_el.getparent().remove(xml_el)


def set_cell_text(cell, text, bold=None, size=None, shade=None, align=None):
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r.text = ""
    run = p.runs[0] if p.runs else p.add_run("")
    run.text = text
    if bold is not None:
        run.bold = bold
    if size is not None:
        run.font.size = Pt(size)
    if align is not None:
        p.alignment = align
    if shade is not None:
        tcPr = cell._tc.get_or_add_tcPr()
        shd = tcPr.find(qn("w:shd"))
        if shd is None:
            shd = tcPr.makeelement(qn("w:shd"), {})
            tcPr.append(shd)
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), shade)


def style_table_borders(table):
    """Apply simple single-line borders to every cell, matching the real
    sample's table style (python-docx has no high-level API for this)."""
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = tblPr.makeelement(qn("w:tblBorders"), {})
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.makeelement(qn(f"w:{edge}"), {})
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "999999")
        borders.append(el)
    tblPr.append(borders)


HEADER_SHADE = "E6E6E6"
SECTION_SHADE = "D9D9D9"
