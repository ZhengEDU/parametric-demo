#!/usr/bin/env python3
"""
Builds the Temp/RH Meter automation-ready template (Demo Type A).

Reuses the real sample's shell (logo, header/footer, fonts, legal
boilerplate, signature block, metadata grid, standards table, and the
Environmental Conditions paragraphs) but replaces the "Calibration Data"
table with an 11-column layout (adds an Uncertainty column vs. the real
Weight Set sample) and two measurement groups (Temperature, Humidity), per
Section 4A of the product brief. Demo-authored — not scanned from a real
Temp/RH certificate. See docs/ASSUMPTIONS.md.
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from docx_common import load_shell, move_after, remove, set_cell_text, style_table_borders, HEADER_SHADE, ROOT
from docx.shared import Pt
from docx.oxml.ns import qn

OUT_DIR = ROOT / "server" / "templates" / "temp-rh-meter"
OUT = OUT_DIR / "source.docx"

COLS = [
    # "Standard Reading" (not bare "Standard") whenever the table has an
    # Uncertainty column — confirmed across every 11-column real sample now
    # on hand (Refrigerator, CO2 Incubator, pH Meter, Temperature
    # Transmitter, Thermocouple, Tachometer, Environmental Chamber — 12+
    # files, 100% consistent); 10-column tables (no Uncertainty) use plain
    # "Standard" instead. See docs/DOCUMENT_ANALYSIS.md.
    "Target Value", "Units", "Standard Reading", "As Found", "Deviation",
    "Standard Reading", "As Left", "Deviation", "Uncertainty", "Cal Tolerance",
    "Adjustment Made",
]


def tag_cell(cell, tag_text, paragraph_index=0):
    p = cell.paragraphs[paragraph_index]
    if not p.runs:
        p.add_run(tag_text)
        return
    p.runs[0].text = tag_text
    for r in p.runs[1:]:
        r.text = ""


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    d = load_shell()

    # Issue Date header line, above the title. Real evidence (12+ of the 31
    # real samples now on hand, spanning Jun-Sep 2026 and 5 different
    # customers) open with an "Issue Date: {date}  ...  Certificate
    # #3865.01" line that the original single Weight Set sample this shell
    # is cloned from does not have. The trailing "Certificate #3865.01" is
    # BYTE-IDENTICAL across every one of those 12+ samples regardless of
    # date or customer — confirmed not to be a per-job/per-visit number, so
    # it's reproduced here as the same static literal text Parametric's own
    # real template prints, rather than computed from anything. Paragraph 0
    # is a blank, centered spacer in the shell; repurposed here instead of
    # inserting a new one. See docs/ASSUMPTIONS.md.
    p_issue = d.paragraphs[0]
    p_issue.alignment = None
    run = p_issue.add_run("Issue Date: ")
    run.font.size = Pt(8)
    run2 = p_issue.add_run("{calibration_date}")
    run2.font.size = Pt(8)
    run3 = p_issue.add_run("\t" * 10)
    run3.font.size = Pt(8)
    run4 = p_issue.add_run("Certificate #3865.01")
    run4.font.size = Pt(8)

    # Environmental Conditions paragraphs (7,8,9) + metadata grid (table 0)
    # + standards table (table 1): identical tagging to the Weight Set
    # template — same helper logic, duplicated inline to keep each build
    # script independently runnable and auditable.
    p_temp = d.paragraphs[8]
    p_temp.runs[4].text = "{env_temperature_value}"
    p_temp.runs[10].text = "{env_temperature_as_found_status}"
    p_temp.runs[11].text = ""
    p_temp.runs[14].text = "{env_temperature_as_left_status}"
    p_temp.runs[21].text = "{env_temperature_status}"
    p_temp.runs[26].text = "{calibration_location}"
    p_temp.runs[27].text = ""

    p_hum = d.paragraphs[9]
    p_hum.runs[4].text = "{env_humidity_value}"
    p_hum.runs[5].text = ""

    p_sig = d.paragraphs[18]
    p_sig.runs[0].text = "{customer_name}"

    t0 = d.tables[0]
    tag_cell(t0.rows[0].cells[1], "{certificate_number}")
    tag_cell(t0.rows[0].cells[3], "{asset_number}")
    tag_cell(t0.rows[1].cells[1], "{customer_name}")
    tag_cell(t0.rows[1].cells[3], "{asset_description}")
    tag_cell(t0.rows[2].cells[1], "{customer_address_line1}")
    tag_cell(t0.rows[2].cells[3], "{asset_manufacturer}")
    tag_cell(t0.rows[3].cells[1], "{customer_address_line2}")
    tag_cell(t0.rows[3].cells[3], "{asset_model}")
    tag_cell(t0.rows[4].cells[1], "{customer_city_state_zip}")
    tag_cell(t0.rows[4].cells[3], "{asset_serial_number}")
    tag_cell(t0.rows[5].cells[1], "{customer_contact}")
    tag_cell(t0.rows[5].cells[3], "{asset_accuracy}")
    tag_cell(t0.rows[6].cells[1], "{customer_phone}")
    tag_cell(t0.rows[6].cells[3], "{asset_range}")
    tag_cell(t0.rows[7].cells[1], "{calibration_date}")
    tag_cell(t0.rows[7].cells[3], "{procedure_name}")
    tag_cell(t0.rows[8].cells[1], "{calibration_next_due}")
    tag_cell(t0.rows[8].cells[3], "{calibration_interval}")

    t1 = d.tables[1]
    tag_cell(t1.rows[1].cells[0], "{#standards}{standard_id_number}")
    tag_cell(t1.rows[1].cells[1], "{standard_manufacturer}")
    tag_cell(t1.rows[1].cells[2], "{standard_model_number}")
    tag_cell(t1.rows[1].cells[3], "{standard_description}")
    tag_cell(t1.rows[1].cells[4], "{standard_cal_due}{/standards}")

    # Comments table (table 3)
    t3 = d.tables[3]
    tag_cell(t3.rows[0].cells[0], "{comments}", paragraph_index=1)

    # --- Rebuild the Calibration Data table with 11 columns ---
    old_cal_table = d.tables[2]
    anchor = old_cal_table
    new_tbl = d.add_table(rows=3, cols=11)
    style_table_borders(new_tbl)
    move_after(anchor, new_tbl)
    remove(old_cal_table)

    for ci, label in enumerate(COLS):
        set_cell_text(new_tbl.rows[0].cells[ci], label, bold=True, size=9, shade=HEADER_SHADE)

    # Row 1: group-label row, 4 merged spans: (0-1)=2, (2-4)=3, (5-7)=3, (8)=1, (9-10)=2
    r1 = new_tbl.rows[1]
    g_mass = r1.cells[0].merge(r1.cells[1])
    g_found = r1.cells[2].merge(r1.cells[4])
    g_left = r1.cells[5].merge(r1.cells[7])
    _ = r1.cells[8]  # uncertainty column, left un-merged/blank on the group row
    g_tail = r1.cells[9].merge(r1.cells[10])
    set_cell_text(g_mass, "{#groups}{group_name}", bold=True, size=9)

    # Row 2: templated point row
    r2 = new_tbl.rows[2]
    set_cell_text(r2.cells[0], "{#points}{point_target_value}", size=9)
    set_cell_text(r2.cells[1], "{point_unit}", size=9)
    set_cell_text(r2.cells[2], "{point_standard_as_found}", size=9)
    set_cell_text(r2.cells[3], "{point_as_found}", size=9)
    set_cell_text(r2.cells[4], "{point_deviation_as_found}", size=9)
    set_cell_text(r2.cells[5], "{point_standard_as_left}", size=9)
    set_cell_text(r2.cells[6], "{point_as_left}", size=9)
    set_cell_text(r2.cells[7], "{point_deviation_as_left}", size=9)
    set_cell_text(r2.cells[8], "{point_uncertainty}", size=9)
    set_cell_text(r2.cells[9], "{point_cal_tolerance}", size=9)
    set_cell_text(r2.cells[10], "{point_adjustment_made}{/points}{/groups}", size=9)

    d.save(str(OUT))
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
